"""Targeted checks for corrected document tables, figures and PDF page bounds."""
import csv
import hashlib
import json
from pathlib import Path
from docx import Document
import pdfplumber

root = Path.cwd()
out = root / 'outputs/submission_review'
main = Document(out / 'Chapter1_manuscript_REVIEW.docx')
supp = Document(out / 'Chapter1_supplement_REVIEW.docx')

def all_text(doc):
    return '\n'.join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for row in t.rows for c in row.cells])

mt, st = all_text(main), all_text(supp)
assert '414' in mt and '442' in mt and '431' in mt
assert '254 inside, 264 outside' in mt
assert '−0.235 [−0.753, 0.284]' in mt
assert 'annual cropland fractions were unavailable' not in mt.lower()
assert 'all three' in st.lower() and '2001–2019' in st
assert 'Katanga Protection Cluster. (2015)' in st and 'Asmani 2015' not in st
assert 'Radio Okapi 2013a' not in st
for num in range(18, 22):
    assert 'Table S' + str(num) in st
assert 'Figure S3.' in st and '93.99%' in st and 'missing' in st
assert len(main.inline_shapes) == 4 and len(supp.inline_shapes) == 3
shape = supp.inline_shapes[0]
part = supp.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed]
assert hashlib.sha256(part.blob).digest() == hashlib.sha256(
    (root / 'outputs/publication/supplement/figures/Figure_S2_threshold_spatial_sensitivities.png').read_bytes()).digest()
# All six 10% corridor fire profile/pairwise rows have retained numerical CIs.
with (root / 'outputs/publication/supplement/tables/Table_S11_threshold_sensitivity.csv').open(encoding='utf-8-sig') as f:
    tab = list(csv.DictReader(f))
fire = [r for r in tab if r.get('Outcome') == 'Fire' and r.get('Threshold') == '10%']
assert len(fire) == 6, (len(fire), tab[0].keys())
assert all(r.get('Lower 95%') not in ['', 'NA'] for r in fire)
for source in json.loads((out / 'source_document_provenance.json').read_text()):
    p = Path(source['path'])
    if p.exists():
        assert hashlib.sha256(p.read_bytes()).hexdigest() == source['sha256']
pdf_checks = []
for p in sorted(out.glob('*.pdf')):
    with pdfplumber.open(p) as pdf:
        for n, page in enumerate(pdf.pages, 1):
            chars = page.chars
            outside = [c for c in chars if c['x0'] < -1 or c['x1'] > page.width + 1
                       or c['top'] < -1 or c['bottom'] > page.height + 1]
            assert chars or page.images, (p, n, 'empty page')
            assert not outside, (p, n, 'out-of-page text')
        pdf_checks.append({'file': p.name, 'pages': len(pdf.pages), 'page_bounds': 'pass'})
assert len(pdf_checks) == 4
(root / 'docs/submission/validation/review-pdf-checks.json').write_text(json.dumps(pdf_checks, indent=2))
print('Review text, tables, figure linkage, unchanged sources and PDF page bounds passed.')
