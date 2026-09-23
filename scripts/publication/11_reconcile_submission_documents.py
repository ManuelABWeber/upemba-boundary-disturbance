"""Build separate review copies; never overwrite source manuscripts."""
import csv, json, hashlib, re, os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path.cwd()
OUT = ROOT / 'outputs/submission_review'
OUT.mkdir(parents=True, exist_ok=True)
SRC = Path(os.environ.get('UPEMBA_MANUSCRIPT_ROOT', 'C:/0_Documents/biological conservation manuscript/Biological_Conservation_submission_package'))
sources = [SRC / 'Biological_Conservation_anonymized_manuscript.docx', SRC / 'Biological_Conservation_supplementary_material.docx']
main, supp = [Document(p) for p in sources]
changes = []

def paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs

def find(doc, prefix):
    hits = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(hits) == 1, (prefix, len(hits))
    return hits[0]

def replace(doc, prefix, text, key):
    p = find(doc, prefix)
    changes.append(dict(id=key, old=p.text, new=text))
    p.text = text
    return p

def note(doc, p, text):
    if not p.runs:
        p.add_run(' ')
    doc.add_comment(p.runs[0], text=text, author='Reconciliation review', initials='RR')

def rows(path):
    return list(csv.DictReader(open(ROOT / path, encoding='utf-8-sig')))

profiles = rows('outputs/final_reconciliation/agriculture_profile_estimates_hc3.csv')
boot = rows('outputs/final_reconciliation/agriculture_spatial_bootstrap_intervals.csv')
counts = rows('outputs/final_reconciliation/agriculture_counts_by_side_group_profile_year.csv')
labels = {'Neither actor dominant': 'Fragmented', 'Militia dominant': 'Militia-centred', 'Park dominant': 'Park-centred'}
names = {'first_crossing_full': 'First crossing, 2001–2022', 'first_crossing_restricted': 'First crossing, 2001–2019', 'persistence_corrected': 'Persistent establishment, 2001–2019', 'annual_presence': 'Annual presence, 2001–2022'}

def ci(row):
    return f"{float(row['estimate']):.3f} [{float(row['conf_low']):.3f}, {float(row['conf_high']):.3f}]"

def total(name, col):
    return sum(int(r[col]) for r in counts if r['analysis'] == name)

replace(main, 'Protected-area evaluations often treat', 'Protected-area evaluations often treat legal designation as a stable intervention, although practical authority can change while boundaries remain fixed. We examined whether Upemba National Park’s legal boundary aligned with disturbance under fragmented, militia-centred, and park-centred territorial-control profiles. We combined a historical chronology with recurrent fire and first tree-cover-loss and agricultural-expansion threshold crossings in 500 m cells during 2001–2022. We estimated profile-specific inside–outside contrasts within landscape groups and compared the legal boundary with ten outward pseudo-boundaries.', 'abstract-design')
replace(main, 'At the primary 25% cumulative-loss threshold, tree-cover-loss odds', 'At the primary threshold, tree-cover-loss odds were 57–76% lower inside under all profiles; pairwise profile differences remained uncertain. Legal-boundary estimates were more negative than all evaluated outward estimates. This result weakened in the better-balanced 5 km corridor and at the sparse 50% threshold. Fire contrasts were higher under fragmented than militia-centred or park-centred control, but comparable contrasts and sign reversals occurred away from the legal line. Agricultural first-expansion estimates remained sensitive to the analysis period, persistence definition and uncertainty calculation.', 'abstract-fire-comparison')
replace(main, 'The legal-boundary-aligned tree-cover-loss pattern across', 'A favourable land-surface outcome did not identify the actor, mechanism, conservation intent or legitimacy that produced it. Historical reconstruction, process-specific satellite observations and boundary-location diagnostics therefore answer distinct questions when protected-area authority changes over time. The estimates describe conditional spatial associations rather than causal effects of legal protection.', 'abstract-conclusion')
replace(main, 'Protected areas are commonly evaluated', 'Legal designation establishes jurisdiction: it defines a protected area’s boundary and formal rules. Practical authority concerns who can implement or contest those rules. Protected areas are commonly evaluated by comparing environmental conditions inside and outside their boundaries, often finding lower habitat loss or pressure inside (Andam et al., 2008; Geldmann et al., 2013, 2019; Gray et al., 2016; McNicol et al., 2023). Yet the capacity to regulate access can change while the legal line remains fixed. Protected areas territorialise space through rules and authorised actors, but these institutions remain contested and historically variable (Bassett and Gautier, 2014; Holmes, 2014). Their environmental expression therefore requires evidence about practical authority as well as legal jurisdiction.', 'introduction-jurisdiction')
replace(main, 'Agricultural expansion used the Annual', 'Agricultural expansion used the Annual Cropland Extent Dataset for Africa (Lou et al., 2025). An event was the first observed transition from below to at least 25% cropland, after which the cell left the risk set. Baseline agricultural cells were retained, so a subsequent decline and recrossing could define their first detected expansion within the observation window. This is not first-ever cultivation. We reconstructed annual 2000–2022 cropland fractions and verified every first-crossing date against the primary input. These observations support separate persistence, reversal and annual-presence analyses (Supplementary Methods S7; Tables S18–S20). None measures net agricultural change.', 'agriculture-availability')
replace(main, 'We varied event threshold', 'We varied event threshold (10%, 25%, 50%), spatial domain (5, 10, 20 km and full domain), and weighting. We distinguished the earlier exploratory chronology checks (Table S12) from the recovered, matched primary-corridor analyses of one- and two-year profile lags and transition-year exclusions (Supplementary Methods S5; Table S21). These assess temporal alignment, with altered eligible years; they do not identify delayed causal effects. Fire harmonisation and optimizer checks are reported in Supplementary Methods S3–S4 and Table S13. The retained 10% and 25% BFGS fits meet strict numerical criteria; the 50% fit does not meet the independent-gradient criterion.', 'sensitivity-consolidation')
find(main, 'We report whether each legal').insert_paragraph_before('The park-facing band at the +15 km pseudo-boundary spans +5–15 km from the legal line and overlaps the primary outside band over +5–10 km. It is not an independent outside control. Overlap among offset corridors also prevents treating their estimates as independent replications.')
anchor = find(main, '3. Results')
anchor.insert_paragraph_before('2.7 Additional agricultural and seasonal diagnostics', 'Heading 2')
anchor.insert_paragraph_before('Persistent first establishment required a crossing to at least 25% cropland and at least two of the next three calendar years at that threshold, including year +1 or +2. All three follow-ups had to be valid. Candidate years were 2001–2019, with follow-up through 2022. We compared this model with first crossing over both 2001–2019 and the original 2001–2022 period, preserving the landscape groups, corridor and two-stage specification. A later qualifying recrossing could replace an earlier transient crossing; cells exited after the dated qualifying crossing, not its confirmation year.')
anchor.insert_paragraph_before('For annual presence, all eligible cells remained in the annual 2001–2022 sample, including baseline agricultural cells and cells observed after establishment. We applied the same paired landscape-year log-odds design to annual classifications and calculated intervals from 499 resamples of whole cell histories in fixed 5 km spatial blocks. Blocks were sampled within landscape groups; boundary-crossing blocks retained both sides. This addresses repeated observations and local spatial dependence conditional on the observed chronology, but not dependence between blocks or uncertain historical coding. We separated two-year classification reversal from persistent establishment, and static Hansen loss–gain overlap from dated recovery (Supplementary Methods S7).')
anchor.insert_paragraph_before('An appendix figure compares monthly CHIRPS v2.0 rainfall with unfiltered monthly FireCCI51 mapped burned area in the same corridor during 2021 (Fig. S3; Supplementary Methods S6). We selected the year using observation coverage, weighted precipitation by area and displayed missing-fire-area bounds. The figure illustrates one seasonal cycle rather than a causal or long-term rainfall–fire relationship.')
replace(main, 'Agricultural-expansion estimates were inconclusive.', f"The primary agricultural model retained 518 events (254 inside, 264 outside). Cell IDs, coordinates, sides, dates and corridor flags agreed between the primary input and post-event classifications. Restriction to 2001–2019 retained {total('first_crossing_restricted', 'y')} first crossings, and corrected persistence retained 414 qualifying establishments. The park-centred contrast changed from −0.347 [−0.657, −0.037] in the full-period first-crossing model to −0.240 [−0.611, 0.131] in its restricted refit, then to −0.235 [−0.753, 0.284] under corrected persistence (Table S18). Shortening the period explained most of the point-estimate shift. The persistence standard error increased; overlapping intervals do not establish equivalent effects. The superseded 431-event estimate, −0.237 [−0.642, 0.168], used incomplete follow-up and is not the corrected result.", 'agriculture-results')
find(main, 'Figure 3.').insert_paragraph_before('Among original corridor crossings with two valid follow-ups, 10/236 inside (4.2%) and 13/227 outside (5.7%) fell below 25% in both following years; 18 inside and 37 outside events lacked sufficient follow-up. Full-domain values were 10/237 and 193/2,092 (Table S19). These are classification reversals, not ecological recovery. Annual-presence estimates were negative under all profiles (−1.549, −1.356 and −1.175); spatial-bootstrap intervals were much wider than HC3 intervals and included zero for militia-centred and park-centred periods (Table S18). Prevalence answers a different question from first expansion, with substantial spatial uncertainty.')
replace(main, 'Figure 3. Primary pooled', 'Figure 3. Primary pooled contrasts at the 25% threshold within the symmetric 10 km corridor. Panel A shows profile-specific inside-minus-outside log odds: positive values indicate higher disturbance inside, negative values lower disturbance inside. Panel B shows the first named profile’s contrast minus the second named profile’s contrast: positive values mean the contrast is more positive under the first profile, not necessarily higher disturbance inside. Bars are 95% conditional intervals calculated from the fitted covariance: SE(Lβ) = √(LΣLᵀ), including covariance between coefficients. They are not differences between interval endpoints. Fire uses a beta-binomial model; sparse outcomes use weighted two-stage models and HC3 covariance. Unresolved pairwise differences do not establish equivalence. Primary intervals do not account for all spatial or temporal dependence.', 'figure3-signs-covariance')
find(main, 'Figure 3.')._p.addnext(find(main, 'Among original corridor crossings')._p)
p = find(main, 'Several limitations constrain')
p.text = p.text.replace('Chronology diagnostics were not directly matched to the primary sparse-outcome model and cannot identify response lags.', 'Recovered matched lag analyses check temporal alignment but cannot identify causal delays. The spatial bootstrap broadens agricultural uncertainty; primary fire and tree-loss intervals remain conditional on unmodelled dependence.').replace('First-event outcomes omit regrowth, reclearing, abandonment, and shifting cultivation.', 'First-event outcomes do not describe subsequent land states. Annual cropland fractions support the separate diagnostics reported here, but threshold reversal does not demonstrate abandonment or recovery. Static Hansen gain cannot establish post-loss recovery (Supplementary Methods S7).')
p = replace(main, 'Data and code availability.', 'Data and code availability. Source satellite products are available from the providers in Tables S1 and S16. Scripts, compact derived tables and reproducibility records are retained in the analytical repository. The repository is currently private; an anonymised reviewer-access route and a persistent public deposit must be confirmed before submission. Restricted historical documents require rights-holder permission and appropriate security safeguards; Table S6 identifies the evidence used.', 'data-access-status')
note(main, p, 'Submission blocker: confirm anonymised reviewer access and final deposit identifier. No public deposit or redistribution permission is asserted.')
p = replace(main, 'During preparation of this work, the authors used ChatGPT', 'During preparation, ChatGPT and Codex (OpenAI) assisted with structural and language revision, code correction, reproducibility checks and generation of data-derived plots. Author review, editing, verification and responsibility for the final article must be confirmed before submission.', 'ai-declaration-draft')
note(main, p, 'Draft disclosure: replace the final sentence only after authors have reviewed the results and accepted responsibility. Do not submit this review wording unchanged.')
note(main, main.paragraphs[0], 'Review copy based on the Biological Conservation submission-package manuscript. The newer commented files and cited S7 passage were not located. Source documents remain untouched. Confirm final version, author details and declarations before submission.')
# Registered article number, and removal of two references orphaned in source.
body = '\n'.join(p.text for p in main.paragraphs).split('\nReferences\n')[0]
for p in list(main.paragraphs):
    if p.text.startswith('McNicol, I.M.'):
        p.text = p.text.replace('4, 1–14.', '4, 392.')
    if any(p.text.startswith(name + ',') and name not in body for name in ['Fromont', 'Gohr']):
        p._element.getparent().remove(p._element)

replace(supp, 'The transition-year and episode-omission diagnostics', 'Table S12 retains earlier exploratory chronology diagnostics, whose sparse-outcome domain and implementation differed from the primary model. Recovered matched primary-corridor lag and transition-exclusion analyses are now reported in Supplementary Methods S5 and Table S21. Neither estimates a causal response delay.', 'supp-chronology')
find(supp, 'The boundary-location diagnostic').text += ' The +15 km park-facing band is +5–15 km and overlaps the primary outside band over +5–10 km.'
anchor = find(supp, 'Supplementary Tables')
for block in (ROOT / 'docs/submission/replacement-supplementary-methods.md').read_text(encoding='utf-8').split('\n\n'):
    block = block.strip()
    if block:
        anchor.insert_paragraph_before(block.removeprefix('## '), 'Heading 2' if block.startswith('## ') else None)

# Match retained refined fire estimates and pairwise orientation in Table S11A.
fire = rows('analysis_fire_optimizer_refinement_dev/tables/fire_10km_optimizer_profile_contrasts.csv')
pairs = rows('analysis_fire_optimizer_refinement_dev/tables/fire_10km_optimizer_pairwise_differences.csv')
def norm(x):
    return re.sub('[^a-z]+', ' ', x.lower().replace(' control', '')).strip()
for row in supp.tables[11].rows[1:]:
    if 'fire' not in row.cells[0].text.lower():
        continue
    label = norm(row.cells[2].text)
    pool = pairs if 'pair' in row.cells[1].text.lower() else fire
    for col, tag in [(3, 'tau010'), (4, 'tau025'), (5, 'tau050')]:
        for r in pool:
            if r['threshold_tag'] != tag or r['optimizer_config'] != 'D_BFGS_diagnostic':
                continue
            human = r.get('profile_comparison', r.get('governance_profile', ''))
            for raw, public in labels.items():
                human = human.replace(raw, public)
            sg = 1 if norm(human) == label else -1 if ' minus ' in human and norm(' minus '.join(human.split(' minus ')[::-1])) == label else 0
            if not sg:
                continue
            est = sg * float(r['estimate'])
            if r['lower_95_ci']:
                lo, hi = sorted([sg * float(r['lower_95_ci']), sg * float(r['upper_95_ci'])])
                row.cells[col].text = f'{est:.3f} [{lo:.3f}, {hi:.3f}]'
            else:
                row.cells[col].text = f'{est:.3f} [non-strict fit]'

for p in paragraphs(supp):
    text = p.text.replace('Parc Upemba: l’ICCN dénonce le braconnage des habitants d’un village dispersé', 'Parc Upemba: l’ICCN accuse des villageois de se livrer au braconnage').replace('Publié le ven, 14/03/2014 - 15:29.', 'Published 14 March 2014; updated 8 August 2015.')
    for yr in ['2010', '2013', '2021', '2024']:
        text = text.replace('Radio Okapi (' + yr + 'a)', 'Radio Okapi (' + yr + ')').replace('Radio Okapi, ' + yr + 'a', 'Radio Okapi, ' + yr).replace('(' + yr + 'a).', '(' + yr + ').')
        text = text.replace('Radio Okapi ' + yr + 'a', 'Radio Okapi ' + yr)
    text = text.replace('Asmani 2015', 'Katanga Protection Cluster 2015')
    if text.startswith('Asmani,'):
        text = 'Katanga Protection Cluster. (2015). Violence, displacement and humanitarian action in Katanga: A report on the protection of civilians in the south-eastern DR Congo province of Katanga. Produced with the support of OCHA, May 2015. https://www.ecoi.net/en/file/local/1095206/1930_1401096604_kantaga-report-2014-final-english-22052014-001.pdf'
    if text != p.text:
        p.text = text
replace(supp, 'Data and code availability.', 'Data and code availability. Satellite products and acquisition links are public. The analytical repository is currently private; authors must confirm anonymised reviewer access and a persistent deposit before submission. Restricted historical evidence is not included in the public-data deliverables. Compact numerical outputs and acquisition/processing manifests accompany this review.', 'supp-data-access')

anchor = find(supp, 'Supplementary Figures')
def table_at(title, headers, data):
    anchor.insert_paragraph_before(title, 'Caption')
    tab = supp.add_table(rows=1, cols=len(headers))
    for cell, text in zip(tab.rows[0].cells, headers):
        cell.text = text
    for vals in data:
        for cell, val in zip(tab.add_row().cells, vals):
            cell.text = str(val)
    anchor._p.addprevious(tab._tbl)
    return tab

data = []
for r in profiles:
    b = next(b for b in boot if b['analysis'] == r['run_id'] and b['governance_profile'] == r['governance_profile'])
    data.append([names[r['run_id']], labels[r['governance_profile']], ci(r), f"[{float(b['conf_low']):.3f}, {float(b['conf_high']):.3f}]"])
table_at('Table S18. Agricultural inside-minus-outside log odds. HC3 intervals retain the established comparison; the final column reports 499 whole-trajectory 5 km spatial-bootstrap percentile intervals. Annual presence is prevalence, a different response from first establishment.', ['Response / period', 'Profile', 'Estimate [HC3 CI]', 'Spatial-bootstrap CI'], data)
revs = rows('outputs/final_reconciliation/agriculture_reversal_denominators.csv')
table_at('Table S19A. Reversal after the original crossing. Denominators require both follow-ups. No missing-data cases occurred.', ['Domain', 'Side', 'All events', 'Reversed / eligible', 'Insufficient follow-up'], [[r['domain'], r['park_side'], r['all_events'], f"{r['reversed']} / {r['eligible']}", r['insufficient_followup']] for r in revs])
table_at('Table S19B. Static gain among loss-threshold crossings through 2012. Neither gain measure establishes post-loss recovery.', ['Domain', 'Crossings', 'Any gain in cell', 'Native overlap'], [['Full retained domain', 2801, 754, 599], ['10 km corridor', 364, 58, 46]])
data = []
for name in names:
    for prof in labels:
        c = [r for r in counts if r['analysis'] == name and r['governance_profile'] == prof]
        data.append([names[name], labels[prof], sum(int(r['y']) for r in c if r['side'] == 'inside'), sum(int(r['y']) for r in c if r['side'] == 'outside'), sum(int(r['n']) for r in c)])
table_at('Table S20. Response support by profile and side. Annual-presence numerators count agricultural cell-years; first-event numerators count events. Denominators are at-risk or observed cell-years. Group-by-year counts and all coefficients accompany the analysis.', ['Response / period', 'Profile', 'Inside', 'Outside', 'Cell-years'], data)
lags = rows('outputs/final_robustness/temporal_lag_profile_estimates.csv')
support = rows('outputs/final_robustness/temporal_lag_support.csv')
data = [[r['outcome'].replace('_', ' '), r['specification'], labels[r['governance_profile']], ci(r), next(z['events'] for z in support if z['run_id'] == r['run_id'])] for r in lags]
table_at('Table S21. Matched primary-corridor temporal-alignment checks. Intervals retain conditional model covariance. Lags change eligible years; transition exclusions remove specified years. These are not causal-delay estimates. Detailed support and diagnostics accompany the CSV outputs.', ['Outcome', 'Alignment', 'Profile', 'Estimate [95% CI]', 'Events'], data)

anchor = find(supp, 'Data and code availability.')
# The first existing illustration is the threshold-sensitivity figure. Replace
# its stale raster with the regenerated plot using retained BFGS fire intervals.
old_shape = supp.inline_shapes[0]
image_part = supp.part.related_parts[old_shape._inline.graphic.graphicData.pic.blipFill.blip.embed]
image_part._blob = (ROOT / 'outputs/publication/supplement/figures/Figure_S2_threshold_spatial_sensitivities.png').read_bytes()
from PIL import Image
with Image.open(ROOT / 'outputs/publication/supplement/figures/Figure_S2_threshold_spatial_sensitivities.png') as im:
    old_shape.height = int(old_shape.width * im.height / im.width)
p = anchor.insert_paragraph_before('')
p.add_run().add_picture(str(ROOT / 'outputs/final_reconciliation/rainfall_fire/Figure_S3_rainfall_fire_2021.png'), width=Inches(6.4))
caption = (ROOT / 'outputs/final_reconciliation/rainfall_fire/figure_caption_and_provenance.md').read_text(encoding='utf-8').split('\n', 1)[1].split('\n\n')[0].strip()
anchor.insert_paragraph_before('Figure S3. Seasonal rainfall and mapped burning in 2021. ' + caption + ' Codex assisted the reproducible plotting code; the plotted values derive from the stated satellite products.', 'Caption')
note(supp, supp.paragraphs[0], 'Review supplement: S5–S7 and Tables S18–S21 contain the verified corrections. The S7 Word document described in the request was unavailable. Reconcile its comments before submission.')
note(main, find(main, 'Ethics statement.'), 'Author confirmation required: verify this inherited statement against all historical evidence and any prior interviews used in the chronology. No ethics determination was made in this reconciliation.')
ref_heading = find(supp, 'Supplementary References')
after_heading = False
reference_paragraphs = []
for paragraph in supp.paragraphs:
    if paragraph._p is ref_heading._p:
        after_heading = True
    elif after_heading and paragraph.text.strip():
        reference_paragraphs.append(paragraph)
previous = ref_heading._p
for paragraph in sorted(reference_paragraphs, key=lambda p: p.text.casefold()):
    previous.addnext(paragraph._p)
    previous = paragraph._p

for doc in [main, supp]:
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.keep_with_next = True
        if p.text.startswith('Figure '):
            p.paragraph_format.keep_with_next = False
    if doc is supp:
        for shape in list(doc.inline_shapes)[:2]:
            ratio = shape.height / shape.width
            shape.width = Inches(6.0)
            shape.height = int(shape.width * ratio)
    for t in doc.tables:
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
            b = OxmlElement('w:' + edge)
            b.set(qn('w:val'), 'single' if edge in ['top', 'bottom', 'insideH'] else 'nil')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:color'), 'AAAAAA')
            borders.append(b)
        pr = t._tbl.tblPr
        for b in list(pr.findall(qn('w:tblBorders'))): pr.remove(b)
        pr.append(borders)
        for row in t.rows:
            for cell in row.cells:
                cp = cell._tc.get_or_add_tcPr()
                for el in list(cp):
                    if el.tag in [qn('w:shd'), qn('w:tcBorders')]: cp.remove(el)
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = False
                    for r in p.runs:
                        r.font.size = Pt(9)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
main.save(OUT / 'Chapter1_manuscript_REVIEW.docx')
supp.save(OUT / 'Chapter1_supplement_REVIEW.docx')

h = Document()
h.add_heading('Highlights', 0)
highlights = ['Tree-cover-loss contrasts aligned with Upemba’s legal boundary across profiles.', 'Fire contrasts recurred away from the legal boundary.', 'Agricultural inference depended on event definition, period and uncertainty.', 'Environmental outcomes did not identify conservation intent or legitimacy.']
assert all(len(x) <= 85 for x in highlights)
for x in highlights:
    h.add_paragraph(x, style='List Bullet')
h.save(OUT / 'Chapter1_highlights.docx')
c = Document()
c.add_heading('Draft cover letter — author completion required', 0)
for text in ['Dear Editors,', 'Please consider “Separating environmental outcomes from conservation attribution under changing territorial control” as a research paper in Biological Conservation.', 'The study combines historical reconstruction with satellite observations of fire, tree-cover loss and agriculture in Upemba National Park. Profile-specific boundary contrasts and outward-offset diagnostics distinguish environmental outcomes from attribution to institutions or conservation intent. This distinction is relevant where practical authority changes while legal boundaries remain fixed.', 'The review package includes corrected persistence follow-up, matched-period agricultural comparisons, one annual-presence sensitivity and a full-cycle seasonal rainfall–fire illustration with explicit observation limits. Primary fire and tree-cover-loss estimates remain unchanged. Agricultural conclusions remain cautious.', 'Before sending, the corresponding author must supply approved author names and order, affiliations and contact details, confirmation of author approval and exclusive submission, funding and sponsor roles, competing interests, permissions and any required ethics declarations, CRediT roles, reviewer access and data-deposit details. This draft does not make declarations on the authors’ behalf.', 'Sincerely,\n[Corresponding author — to be confirmed]']:
    c.add_paragraph(text)
c.save(OUT / 'Chapter1_cover_letter_DRAFT.docx')
provenance = [{'path': str(p), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest(), 'role': 'editorial source; retained in place'} for p in sources]
(OUT / 'source_document_provenance.json').write_text(json.dumps(provenance, indent=2), encoding='utf-8')
(ROOT / 'docs/submission/manuscript_replacements.json').write_text(json.dumps(changes, ensure_ascii=False, indent=2), encoding='utf-8')
for name, doc in [('manuscript', main), ('supplement', supp)]:
    (OUT / (name + '_review_text.md')).write_text('\n'.join(p.text for p in paragraphs(doc)), encoding='utf-8')
print('Manuscript words:', len(' '.join(p.text for p in main.paragraphs).split()), 'tables:', len(main.tables), 'images:', len(main.inline_shapes))
print('Abstract words:', sum(len(main.paragraphs[i].text.split()) for i in [2, 3, 4]))
